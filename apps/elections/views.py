import io
import uuid
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.db.models import Count, Prefetch
from django.utils import timezone
from django.http import HttpResponse
from rest_framework import generics, status, permissions
from rest_framework.exceptions import APIException
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import OrderingFilter

from apps.core.filters import MinLengthSearchFilter

from .models import Election
from apps.core.cache_invalidation import invalidate_election

from .aggregates import election_results, election_turnout, eligible_voters_count, results_are_visible
from .services import (
    ElectionStateError,
    cancel_election,
    delete_election,
    finish_election,
    start_election,
)
from .serializers import (
    ElectionSerializer, ElectionStudentSerializer,
    TurnoutSerializer, ElectionResultsSerializer
)
from apps.candidates.models import Candidate
from apps.students.models import Student
from apps.voting.models import VoteRecord, Ballot
from apps.core.permissions import IsAdminUserWithRole, IsStudentAuthenticated, IsNotObserver, IsSuperAdmin
from apps.accounts.models import AdminActionLog

class ElectionOperationDenied(APIException):
    """Переводит ElectionStateError в ответ DRF.

    Нужно именно исключение: DRF игнорирует значение, возвращённое из
    perform_destroy и других perform_*-хуков.
    """
    status_code = status.HTTP_400_BAD_REQUEST

    def __init__(self, error):
        self.detail = error.message
        self.default_code = error.code
        super().__init__(error.message, code=error.code)


def get_client_ip(request):
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0]
    return request.META.get('REMOTE_ADDR')

# --- Admin Election Views ---

class AdminElectionListCreateView(generics.ListCreateAPIView):
    permission_classes = [IsAdminUserWithRole]
    serializer_class = ElectionSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    filter_backends = [DjangoFilterBackend, MinLengthSearchFilter, OrderingFilter]
    filterset_fields = ['status', 'university', 'is_featured']
    search_fields = ['title', 'title_ky', 'description']
    ordering_fields = ['starts_at', 'created_at', 'title', 'featured_order']
    ordering = ['-created_at']

    def get_queryset(self):
        user = self.request.user
        # ТЗ п.26: число кандидатов считает база одной агрегацией, а не
        # отдельный COUNT на каждый объект списка. Аннотация дешевле prefetch:
        # для счётчика не нужно тянуть по сети сами строки кандидатов.
        qs = Election.objects.select_related('university', 'created_by').annotate(
            candidates_count_annotated=Count('candidates', distinct=True)
        )
        if getattr(user, 'role', None) != 'super_admin' and not user.is_superuser:
            qs = qs.filter(university_id=user.university_id)
        return qs

    def perform_create(self, serializer):
        user = self.request.user
        university = serializer.validated_data.get('university')
        if getattr(user, 'role', None) != 'super_admin' and not user.is_superuser:
            university = user.university
        election = serializer.save(created_by=user, university=university)
        AdminActionLog.objects.create(
            admin=user,
            action="create_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title, "university": election.university.code},
            ip_address=get_client_ip(self.request)
        )

class AdminElectionDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAdminUserWithRole]
    serializer_class = ElectionSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_queryset(self):
        user = self.request.user
        qs = Election.objects.select_related('university', 'created_by').annotate(
            candidates_count_annotated=Count('candidates', distinct=True)
        )
        if getattr(user, 'role', None) != 'super_admin' and not user.is_superuser:
            qs = qs.filter(university_id=user.university_id)
        return qs

    def perform_update(self, serializer):
        election = serializer.save()
        invalidate_election(election.id)
        AdminActionLog.objects.create(
            admin=self.request.user,
            action="update_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title, "status": election.status},
            ip_address=get_client_ip(self.request)
        )

    def perform_destroy(self, instance):
        # D-02: раньше здесь был `return Response(...)`. DRF возвращаемое значение
        # perform_destroy ИГНОРИРУЕТ и всегда отдаёт 204 — клиенту сообщалось об
        # удалении, которого не было. Ошибку нужно поднимать, а не возвращать.
        title = instance.title
        election_id = str(instance.id)
        try:
            delete_election(instance)
        except ElectionStateError as exc:
            raise ElectionOperationDenied(exc)

        AdminActionLog.objects.create(
            admin=self.request.user,
            action="delete_election",
            target_type="election",
            target_id=election_id,
            details={"title": title},
            ip_address=get_client_ip(self.request)
        )

class AdminElectionStartView(APIView):
    permission_classes = [IsAdminUserWithRole, IsNotObserver]

    def post(self, request, pk):
        try:
            election = Election.objects.get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        try:
            election = start_election(election)
        except ElectionStateError as exc:
            return Response({"error": {"code": exc.code, "message": exc.message}}, status=status.HTTP_400_BAD_REQUEST)

        AdminActionLog.objects.create(
            admin=request.user,
            action="start_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title},
            ip_address=get_client_ip(request)
        )

        return Response({"success": True, "message": "Выборы успешно запущены", "status": election.status})

class AdminElectionFinishView(APIView):
    permission_classes = [IsAdminUserWithRole, IsNotObserver]

    def post(self, request, pk):
        try:
            election = Election.objects.get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        try:
            election = finish_election(election)
        except ElectionStateError as exc:
            return Response({"error": {"code": exc.code, "message": exc.message}}, status=status.HTTP_400_BAD_REQUEST)

        AdminActionLog.objects.create(
            admin=request.user,
            action="finish_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title},
            ip_address=get_client_ip(request)
        )

        return Response({"success": True, "message": "Выборы успешно завершены", "status": election.status})

class AdminElectionCancelView(APIView):
    permission_classes = [IsAdminUserWithRole, IsNotObserver]

    def post(self, request, pk):
        try:
            election = Election.objects.get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        try:
            election = cancel_election(election)
        except ElectionStateError as exc:
            return Response({"error": {"code": exc.code, "message": exc.message}}, status=status.HTTP_400_BAD_REQUEST)

        AdminActionLog.objects.create(
            admin=request.user,
            action="cancel_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title},
            ip_address=get_client_ip(request)
        )

        return Response({"success": True, "message": "Выборы отменены", "status": election.status})

class AdminElectionTurnoutView(APIView):
    permission_classes = [IsAdminUserWithRole]

    def get(self, request, pk):
        try:
            election = Election.objects.select_related('university').get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        turnout = election_turnout(election)

        return Response({
            "election_id": str(election.id),
            "election_title": election.title,
            "status": election.status,
            "starts_at": election.starts_at,
            "ends_at": election.ends_at,
            "total_eligible": turnout["total_eligible"],
            "total_voted": turnout["total_voted"],
            "turnout_percent": turnout["turnout_percent"],
        }, status=status.HTTP_200_OK)

class AdminElectionResultsView(APIView):
    permission_classes = [IsAdminUserWithRole]

    def get(self, request, pk):
        try:
            election = Election.objects.select_related('university').get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        # In accordance with Section 4.5 & 6.2: results are hidden until finished, unless explicitly configured
        # Правило видимости живёт в одном месте — иначе две вьюхи разойдутся
        if not results_are_visible(election):
            return Response({
                "error": {
                    "code": "results_hidden",
                    "message": "Результаты голосования скрыты до официального завершения выборов для соблюдения тайны и чистоты голосования"
                }
            }, status=status.HTTP_403_FORBIDDEN)

        summary = election_results(election)

        return Response({
            "election_id": str(election.id),
            "election_title": election.title,
            "university_name": election.university.name,
            "total_eligible": summary["total_eligible"],
            "total_voted": summary["total_voted"],
            "turnout_percent": summary["turnout_percent"],
            "candidates": summary["candidates"],
        }, status=status.HTTP_200_OK)

class AdminElectionResultsExportView(APIView):
    permission_classes = [IsAdminUserWithRole]

    def post(self, request, pk):
        try:
            election = Election.objects.select_related('university').get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if getattr(request.user, 'role', None) != 'super_admin' and not request.user.is_superuser:
            if election.university_id != request.user.university_id:
                return Response({"error": {"code": "forbidden", "message": "Доступ ограничен вашим университетом"}}, status=status.HTTP_403_FORBIDDEN)

        if not results_are_visible(election):
            return Response({"error": {"code": "results_hidden", "message": "Экспорт доступен только после завершения выборов"}}, status=status.HTTP_403_FORBIDDEN)

        # Тот же расчёт, что и в results/: одна функция вместо двух копий
        summary = election_results(election)
        total_eligible = summary["total_eligible"]
        total_voted = summary["total_voted"]
        turnout_percent = summary["turnout_percent"]

        # Build styled Excel document
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Итоги голосования"

        # Fonts & styling
        header_font = Font(name="Arial", size=14, bold=True, color="011C42")
        sub_font = Font(name="Arial", size=10, italic=True)
        table_header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        table_header_fill = PatternFill(start_color="011C42", end_color="011C42", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='C0C1C6'),
            right=Side(style='thin', color='C0C1C6'),
            top=Side(style='thin', color='C0C1C6'),
            bottom=Side(style='thin', color='C0C1C6')
        )

        ws.append(["ОФИЦИАЛЬНЫЙ ПРОТОКОЛ ИТОГОВ ГОЛОСОВАНИЯ"])
        ws.append([f"Выборы: {election.title}"])
        ws.append([f"Университет: {election.university.name}"])
        ws.append([f"Дата выгрузки: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        ws.append([])
        ws.append(["Показатель", "Значение"])
        ws.append(["Всего избирателей (студентов):", total_eligible])
        ws.append(["Всего проголосовало:", total_voted])
        ws.append(["Итоговая явка (%):", f"{turnout_percent}%"])
        ws.append([])
        ws.append(["Место", "Кандидат", "Факультет", "Курс", "Должность", "Голосов", "Процент"])

        # Строки кандидатов берутся из уже посчитанного агрегата
        for rank, row in enumerate(summary["candidates"], start=1):
            ws.append([
                rank, row["full_name"], row["faculty"], row["course"],
                row["position"], row["votes"], f"{row['percent']}%",
            ])

        # Column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="results_{election.id}.xlsx"'
        return response

# --- Student Available Elections Views ---

class StudentAvailableElectionsView(APIView):
    permission_classes = [IsStudentAuthenticated]

    def get(self, request):
        # Принципал вместо полной модели: экономит SELECT на горячем endpoint'е
        student_id = request.user.id
        university_id = request.user.university_id
        now = timezone.now()

        # Find elections where student has voted
        voted_election_ids = set(
            VoteRecord.objects.filter(student_id=student_id).values_list('election_id', flat=True)
        )

        include_all = request.query_params.get('all', 'false').lower() == 'true'

        # ElectionStudentSerializer разворачивает вложенных кандидатов, а
        # CandidatePublicSerializer читает candidate.university.name и
        # candidate.election.title. Без select_related внутри Prefetch это
        # два лишних запроса НА КАЖДОГО кандидата (ТЗ п.25).
        base = Election.objects.select_related('university').prefetch_related(
            Prefetch(
                'candidates',
                queryset=Candidate.objects.select_related('university', 'election')
                .order_by('order', 'created_at'),
            )
        )

        if include_all:
            elections = base.filter(
                university_id=university_id
            ).exclude(status=Election.Status.DRAFT).order_by('-created_at')
        else:
            elections = base.filter(
                university_id=university_id,
                status=Election.Status.ACTIVE,
                starts_at__lte=now,
                ends_at__gte=now
            ).exclude(id__in=voted_election_ids).order_by('ends_at')

        serializer = ElectionStudentSerializer(elections, many=True, context={'request': request})
        data = serializer.data
        for item in data:
            try:
                item_uuid = uuid.UUID(str(item['id']))
                item['has_voted'] = item_uuid in voted_election_ids
            except (ValueError, TypeError):
                item['has_voted'] = False

        return Response(data, status=status.HTTP_200_OK)

class StudentElectionDetailView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request, pk):
        try:
            election = Election.objects.select_related('university').prefetch_related(
                Prefetch(
                    'candidates',
                    queryset=Candidate.objects.select_related('university', 'election')
                    .order_by('order', 'created_at'),
                )
            ).get(id=pk)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        serializer = ElectionStudentSerializer(election, context={'request': request})
        data = serializer.data

        # If student is authenticated, check if student belongs to this uni and whether they voted
        # Только идентификаторы — без загрузки полной модели студента
        is_student = getattr(request.user, 'is_student', False)
        if is_student:
            has_voted = VoteRecord.objects.filter(
                election=election, student_id=request.user.id
            ).exists()
            data['has_voted'] = has_voted
            data['is_eligible'] = (str(request.user.university_id) == str(election.university_id))
        else:
            data['has_voted'] = False
            data['is_eligible'] = None

        return Response(data, status=status.HTTP_200_OK)

class PublicRecentElectionsView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        # ElectionSerializer рендерит только candidates_count, сами кандидаты
        # не нужны — поэтому аннотация, а не prefetch_related.
        base = Election.objects.select_related('university', 'created_by').annotate(
            candidates_count_annotated=Count('candidates', distinct=True)
        )
        featured = list(base.filter(
            is_featured=True
        ).exclude(status=Election.Status.CANCELLED).order_by('featured_order', '-created_at')[:6])

        if len(featured) >= 3:
            elections = featured
        else:
            featured_ids = [e.id for e in featured]
            remaining_limit = 6 - len(featured)
            remaining = list(base.exclude(
                id__in=featured_ids
            ).exclude(status=Election.Status.CANCELLED).order_by('-created_at')[:remaining_limit])
            elections = featured + remaining

        serializer = ElectionSerializer(elections, many=True, context={'request': request})
        return Response(serializer.data, status=status.HTTP_200_OK)

class AdminFeaturedElectionsManageView(APIView):
    permission_classes = [IsSuperAdmin]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request):
        """List all elections with featured info for superadmin management."""
        elections = Election.objects.select_related('university', 'created_by').annotate(
            candidates_count_annotated=Count('candidates', distinct=True)
        ).exclude(
            status=Election.Status.CANCELLED
        ).order_by('-is_featured', 'featured_order', '-created_at')
        serializer = ElectionSerializer(elections, many=True, context={'request': request})
        return Response(serializer.data, status=status.HTTP_200_OK)

    def patch(self, request, pk=None):
        """Toggle or update featured status of an election."""
        election_id = pk or request.data.get('id')
        if not election_id:
            return Response({"error": {"code": "missing_id", "message": "ID выборов обязателен"}}, status=status.HTTP_400_BAD_REQUEST)

        try:
            election = Election.objects.get(id=election_id)
        except Election.DoesNotExist:
            return Response({"error": {"code": "not_found", "message": "Выборы не найдены"}}, status=status.HTTP_404_NOT_FOUND)

        if 'is_featured' in request.data:
            val = request.data.get('is_featured')
            election.is_featured = str(val).lower() in ['true', '1', 'yes'] if isinstance(val, str) else bool(val)

        if 'featured_order' in request.data:
            try:
                election.featured_order = int(request.data.get('featured_order'))
            except (ValueError, TypeError):
                pass

        if 'cover_image' in request.FILES:
            election.cover_image = request.FILES['cover_image']

        if 'cover_image_url' in request.data:
            election.cover_image_url = request.data.get('cover_image_url', '')

        election.save()
        invalidate_election(election.id)

        AdminActionLog.objects.create(
            admin=request.user,
            action="update_featured_election",
            target_type="election",
            target_id=str(election.id),
            details={"title": election.title, "is_featured": election.is_featured, "order": election.featured_order},
            ip_address=get_client_ip(request)
        )

        serializer = ElectionSerializer(election, context={'request': request})
        return Response(serializer.data, status=status.HTTP_200_OK)

